"""Verify every existing excelClientId in DB matches Excel by name."""
import os, re, pymysql
from openpyxl import load_workbook

XLSX = "/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx"

url = os.environ['DATABASE_URL']
m = re.match(r'mysql://([^:]+):([^@]+)@([^:]+):(\d+)/(\w+)', url)
u, p, h, port, db = m.groups()
c = pymysql.connect(host=h, port=int(port), user=u, password=p, database=db, ssl={'ssl': {}})
cur = c.cursor()
cur.execute("SELECT id, companyName, excelClientId FROM clients WHERE excelClientId IS NOT NULL ORDER BY id")
db_rows = cur.fetchall()
c.close()

wb = load_workbook(XLSX, data_only=True, read_only=True)
ws = wb["Dim_Clients"]
excel = {}
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    if not row or row[0] is None: continue
    try:
        excel[int(row[0])] = str(row[1]).strip()
    except (ValueError, TypeError):
        pass

print("Verifying DB.excelClientId -> Excel name match:\n")
ok = bad = 0
for db_id, name, eid in db_rows:
    excel_name = excel.get(eid, "<missing>")
    match = name.strip().lower() == excel_name.strip().lower()
    flag = "OK " if match else "BAD"
    if match: ok += 1
    else: bad += 1
    if not match:
        print(f"  {flag}  DB:{db_id:4d}  Excel:{eid:4d}")
        print(f"        DB name:    {name}")
        print(f"        Excel name: {excel_name}")
print(f"\nResult: {ok} correct, {bad} incorrect (out of {len(db_rows)})")
