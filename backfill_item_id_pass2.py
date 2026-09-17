"""Second-pass backfill: for transactions still missing itemId,
fall back to description-only lookup against Fact_Sales (most-common Item_ID per description)."""
import os, re
from collections import Counter
from openpyxl import load_workbook
import pymysql

XLSX = "/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx"
url = os.environ["DATABASE_URL"]
m = re.match(r"mysql://([^:]+):([^@]+)@([^:]+):(\d+)/(\w+)", url)
u, p, h, port, db = m.groups()

wb = load_workbook(XLSX, data_only=True, read_only=True)
ws = wb["Fact_Sales"]
rows = list(ws.iter_rows(values_only=True))
hdr = rows[0]
i_desc = hdr.index("Item_Description")
i_id = hdr.index("Item_ID")

# desc -> Counter(item_id)
by_desc: dict[str, Counter] = {}
for r in rows[1:]:
    if r[i_id] is None or r[i_desc] is None:
        continue
    by_desc.setdefault(str(r[i_desc]).strip(), Counter())[str(r[i_id]).strip()] += 1

print(f"Built description->itemId index: {len(by_desc)} distinct descriptions")

c = pymysql.connect(host=h, port=int(port), user=u, password=p, database=db, ssl={"ssl": {}})
cur = c.cursor(pymysql.cursors.DictCursor)
cur.execute("SELECT id, description FROM transactions WHERE type='purchase' AND itemId IS NULL")
missing = cur.fetchall()
print(f"Missing itemId: {len(missing)} rows")

updates = []
for tx in missing:
    desc = (tx["description"] or "").strip()
    counter = by_desc.get(desc)
    if not counter:
        continue
    best_id = counter.most_common(1)[0][0]
    updates.append((best_id, tx["id"]))

cur2 = c.cursor()
cur2.executemany("UPDATE transactions SET itemId=%s WHERE id=%s", updates)
c.commit()
print(f"Pass-2 applied: {len(updates)} updates")

cur.execute("SELECT COUNT(*) AS n FROM transactions WHERE type='purchase' AND itemId IS NOT NULL")
print("with itemId now:", cur.fetchone())
cur.execute("SELECT COUNT(*) AS n FROM transactions WHERE type='purchase' AND itemId IS NULL")
print("still NULL:", cur.fetchone())
