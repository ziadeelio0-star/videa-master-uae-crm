from openpyxl import load_workbook

wb = load_workbook("/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx", data_only=True, read_only=True)
ws = wb["Dim_Clients"]
print(f"{'ID':>4}  {'Name'}")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    if not row or row[0] is None: continue
    name = str(row[1]).strip().lower()
    if any(k in name for k in ['shandgha', 'safa', 'westream']):
        print(f"{row[0]:>4}  {row[1]}")

print("\nAll names containing 'wood' or 'kitchen':")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    if not row or row[0] is None: continue
    name = str(row[1]).strip().lower()
    if 'wood' in name or 'kitchen' in name:
        print(f"  Excel:{row[0]:>3}  {row[1]}")
