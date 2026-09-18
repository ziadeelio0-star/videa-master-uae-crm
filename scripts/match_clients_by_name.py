"""Try to match NULL-excelClientId clients to Excel by name."""
import os, re, pymysql
from openpyxl import load_workbook

XLSX = "/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx"

url = os.environ['DATABASE_URL']
m = re.match(r'mysql://([^:]+):([^@]+)@([^:]+):(\d+)/(\w+)', url)
u, p, h, port, db = m.groups()
c = pymysql.connect(host=h, port=int(port), user=u, password=p, database=db, ssl={'ssl': {}})
cur = c.cursor()
cur.execute("SELECT id, companyName, excelClientId FROM clients ORDER BY id")
db_clients = cur.fetchall()
c.close()

wb = load_workbook(XLSX, data_only=True, read_only=True)
print("Sheets:", wb.sheetnames)

# Find Dim_Clients sheet
dim_sheet = None
for s in wb.sheetnames:
    if 'client' in s.lower() and 'dim' in s.lower():
        dim_sheet = s
        break
if not dim_sheet:
    for s in wb.sheetnames:
        if 'client' in s.lower():
            dim_sheet = s
            break
print(f"Using sheet: {dim_sheet}")

ws = wb[dim_sheet]
header = None
excel_clients = {}  # excel_id -> name
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        header = [str(c).strip().lower() if c else '' for c in row]
        print("Header:", header)
        continue
    if not row or row[0] is None:
        continue
    try:
        cid_idx = next(j for j, h in enumerate(header) if 'id' in h)
        name_idx = next(j for j, h in enumerate(header) if 'name' in h or 'company' in h)
    except StopIteration:
        cid_idx, name_idx = 0, 1
    cid = row[cid_idx]
    name = row[name_idx]
    if cid is not None and name:
        try:
            excel_clients[int(cid)] = str(name).strip()
        except (ValueError, TypeError):
            pass

print(f"\nExcel has {len(excel_clients)} clients (IDs {min(excel_clients)}..{max(excel_clients)})")

def normalize(s):
    return re.sub(r'[^a-z0-9]', '', (s or '').lower())

excel_norm = {normalize(name): (eid, name) for eid, name in excel_clients.items()}

null_clients = [r for r in db_clients if r[2] is None]
print(f"\nMatching {len(null_clients)} NULL-excelClientId DB clients to Excel by normalized name:")
matched = []
unmatched = []
for db_id, name, _ in null_clients:
    norm = normalize(name)
    if norm in excel_norm:
        eid, ename = excel_norm[norm]
        matched.append((db_id, name, eid, ename))
    else:
        unmatched.append((db_id, name))

print(f"  Matched by name: {len(matched)}")
for m in matched[:20]:
    print(f"    DB:{m[0]:4d}  Excel:{m[2]:4d}  {m[1][:50]}")
print(f"  Still unmatched: {len(unmatched)}")
for u in unmatched[:30]:
    print(f"    DB:{u[0]:4d}  {u[1][:60]}")

# Show Excel clients not in DB
db_norm = {normalize(name): (db_id, name) for db_id, name, _ in db_clients}
excel_only = [(eid, n) for n, (eid, n2) in excel_norm.items() for n in [n2] if normalize(n) not in db_norm]
print(f"\nExcel clients not in DB: {len(excel_only)}")
for eo in excel_only[:20]:
    print(f"    Excel:{eo[0]:4d}  {eo[1][:60]}")
