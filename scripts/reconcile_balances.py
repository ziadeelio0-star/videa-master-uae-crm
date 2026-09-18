"""Reconcile clients.outstandingBalance against Excel Client_Balances."""
import os, re, pymysql
from openpyxl import load_workbook

XLSX = "/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx"

url = os.environ['DATABASE_URL']
m = re.match(r'mysql://([^:]+):([^@]+)@([^:]+):(\d+)/(\w+)', url)
u, p, h, port, db = m.groups()
c = pymysql.connect(host=h, port=int(port), user=u, password=p, database=db, ssl={'ssl': {}})
cur = c.cursor()

# Excel balances
wb = load_workbook(XLSX, data_only=True, read_only=True)
ws = wb["Client_Balances"]
header = None
excel_bal = {}
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        header = [str(x).strip().lower() if x else '' for x in row]
        print(f"Client_Balances header: {header}")
        continue
    if not row or row[0] is None: continue
    try:
        cid = int(row[0])
    except (ValueError, TypeError):
        continue
    # find the outstanding column
    out_idx = next((j for j, h in enumerate(header)
                    if 'outstanding' in h or 'balance' in h), None)
    if out_idx is None: out_idx = 3  # heuristic fallback
    bal = row[out_idx]
    if bal is None: bal = 0
    try:
        excel_bal[cid] = float(bal)
    except (ValueError, TypeError):
        excel_bal[cid] = 0.0

# DB balances
cur.execute("SELECT id, companyName, outstandingBalance FROM clients ORDER BY id")
db_bal = {row[0]: (row[1], float(row[2])) for row in cur.fetchall()}

print(f"\nExcel Client_Balances rows: {len(excel_bal)}")
print(f"DB clients: {len(db_bal)}")

# Compare
mismatches = []
for cid in sorted(set(excel_bal) | set(db_bal)):
    excel = excel_bal.get(cid, 0.0)
    db_name, db_amount = db_bal.get(cid, ("<MISSING>", 0.0))
    diff = round(db_amount - excel, 2)
    if abs(diff) > 0.01:
        mismatches.append((cid, db_name, db_amount, excel, diff))

print(f"\nMismatches: {len(mismatches)}")
for m in mismatches[:20]:
    print(f"  ID:{m[0]:3d}  DB:AED{m[2]:>10.2f}  Excel:AED{m[3]:>10.2f}  diff:{m[4]:>+10.2f}  {m[1][:40]}")

print(f"\nTotal DB outstanding:    AED {sum(v[1] for v in db_bal.values()):>12,.2f}")
print(f"Total Excel outstanding: AED {sum(excel_bal.values()):>12,.2f}")
c.close()
