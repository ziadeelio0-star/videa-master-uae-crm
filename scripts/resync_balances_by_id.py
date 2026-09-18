"""Force every client's outstandingBalance to match Excel Client_Balances using id == excel client_id."""
import os, re, sys, pymysql
from openpyxl import load_workbook

XLSX = "/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx"

url = os.environ['DATABASE_URL']
m = re.match(r'mysql://([^:]+):([^@]+)@([^:]+):(\d+)/(\w+)', url)
u, p, h, port, db = m.groups()
c = pymysql.connect(host=h, port=int(port), user=u, password=p, database=db, ssl={'ssl': {}}, autocommit=False)
cur = c.cursor()

wb = load_workbook(XLSX, data_only=True, read_only=True)
ws = wb["Client_Balances"]
balances = {}
oldest = {}
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    if not row or row[0] is None: continue
    try: cid = int(row[0])
    except: continue
    bal = float(row[5] or 0)
    od = row[2]
    balances[cid] = bal
    oldest[cid] = od

# Reset all balances to 0 first, then apply the Excel sheet
cur.execute("UPDATE clients SET outstandingBalance = 0, oldestUnpaidDate = NULL")
for cid, bal in balances.items():
    od = oldest.get(cid)
    if od:
        cur.execute("UPDATE clients SET outstandingBalance=%s, oldestUnpaidDate=%s WHERE id=%s",
                    (bal, od, cid))
    else:
        cur.execute("UPDATE clients SET outstandingBalance=%s WHERE id=%s", (bal, cid))

# Verify totals
cur.execute("SELECT COALESCE(SUM(outstandingBalance),0) FROM clients")
db_total = float(cur.fetchone()[0])
xl_total = sum(balances.values())
print(f"DB total after sync:    AED {db_total:,.2f}")
print(f"Excel total:            AED {xl_total:,.2f}")
print(f"Difference:             AED {db_total - xl_total:,.2f}")

if abs(db_total - xl_total) < 0.01:
    c.commit()
    print("COMMITTED.")
else:
    c.rollback()
    print("ROLLED BACK.")
    sys.exit(1)

c.close()
