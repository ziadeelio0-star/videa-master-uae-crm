#!/usr/bin/env python3
"""Audit Fact_Cost sheet: list every distinct cost_type with totals + sample rows."""
import openpyxl
import sys
from collections import defaultdict

WB = "/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx"

wb = openpyxl.load_workbook(WB, data_only=True)
print("Sheets:", wb.sheetnames)

# Find the cost sheet
candidates = [s for s in wb.sheetnames if "cost" in s.lower()]
print("Cost-related sheets:", candidates)

for sheet_name in candidates:
    ws = wb[sheet_name]
    print(f"\n==== {sheet_name} ({ws.max_row} rows x {ws.max_column} cols) ====")
    headers = [c.value for c in ws[1]]
    print("Headers:", headers)
    # Find cost_type and amount columns
    cost_type_idx = None
    amount_idx = None
    date_idx = None
    for i, h in enumerate(headers):
        if h is None:
            continue
        h_low = str(h).strip().lower()
        if h_low in ("cost_type", "cost type", "type"):
            cost_type_idx = i
        if h_low in ("amount", "cost", "value", "total"):
            if amount_idx is None:
                amount_idx = i
        if h_low in ("date", "transaction_date", "tx_date"):
            date_idx = i
    print(f"cost_type col = {cost_type_idx}, amount col = {amount_idx}, date col = {date_idx}")

    if cost_type_idx is None:
        # Print first 5 rows to figure out
        for r in range(1, min(6, ws.max_row + 1)):
            print([c.value for c in ws[r]])
        continue

    totals = defaultdict(lambda: {"sum": 0.0, "count": 0, "samples": []})
    for r in range(2, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        if all(v is None for v in row):
            continue
        ct = row[cost_type_idx]
        if ct is None:
            ct = "(blank)"
        ct_key = str(ct).strip()
        amt = row[amount_idx] if amount_idx is not None else None
        try:
            amt_f = float(amt) if amt is not None else 0.0
        except Exception:
            amt_f = 0.0
        totals[ct_key]["sum"] += amt_f
        totals[ct_key]["count"] += 1
        if len(totals[ct_key]["samples"]) < 2:
            totals[ct_key]["samples"].append(row)

    print(f"\nDistinct cost_type values ({len(totals)}):")
    for k in sorted(totals, key=lambda x: -totals[x]["sum"]):
        v = totals[k]
        print(f"  {k!r:50s}  count={v['count']:5d}  total={v['sum']:>14,.2f}")

    print("\nSample rows per cost_type (up to 10 types):")
    for k in sorted(totals, key=lambda x: -totals[x]["sum"])[:10]:
        print(f"\n  >>> {k!r}")
        for s in totals[k]["samples"]:
            print("     ", s)
