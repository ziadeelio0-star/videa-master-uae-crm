#!/usr/bin/env python3
"""
Auto-sync script for Cutting Tools CRM
Imports data from Excel dashboard and syncs with database
Supports fuzzy matching for client names
"""

import openpyxl
import mysql.connector
import os
import sys
from datetime import datetime
from difflib import SequenceMatcher

def get_db_connection():
    db_url = os.environ.get('DATABASE_URL', '')
    if not db_url:
        print("ERROR: DATABASE_URL not set")
        sys.exit(1)
    
    db_url = db_url.replace('mysql://', '')
    user_pass, rest = db_url.split('@')
    user, password = user_pass.split(':')
    host_port, database = rest.split('/')
    if '?' in database:
        database = database.split('?')[0]
    if ':' in host_port:
        host, port = host_port.split(':')
        port = int(port)
    else:
        host = host_port
        port = 3306
    
    return mysql.connector.connect(
        host=host, port=port, user=user, password=password, 
        database=database, ssl_disabled=False
    )

def fuzzy_match(name1, name2, threshold=0.75):
    """Check if two names match with fuzzy matching"""
    ratio = SequenceMatcher(None, name1.lower(), name2.lower()).ratio()
    return ratio >= threshold

def find_client_by_name(cursor, client_name):
    """Find client ID by name with fuzzy matching"""
    cursor.execute("SELECT id, companyName FROM clients")
    clients = cursor.fetchall()
    
    for client_id, db_name in clients:
        if fuzzy_match(client_name, db_name):
            return client_id
    return None

def parse_date(cell_value):
    """Parse date from Excel cell"""
    if not cell_value:
        return None
    if isinstance(cell_value, datetime):
        return cell_value
    try:
        if isinstance(cell_value, str):
            # Try DD/MM/YYYY format
            return datetime.strptime(cell_value, '%d/%m/%Y')
    except:
        pass
    return None

def sync_client_balances(cursor, excel_file):
    """Sync Client_Balances sheet with clients table"""
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    
    if 'Client_Balances' not in wb.sheetnames:
        print("WARNING: Client_Balances sheet not found")
        return 0
    
    ws = wb['Client_Balances']
    
    # Find column indices
    header_row = ws[1]
    col_map = {}
    for col_idx, cell in enumerate(header_row, 1):
        if cell.value:
            col_map[cell.value.lower()] = col_idx
    
    required_cols = ['client_name', 'outstanding_balance', 'oldest_unpaid_invoice_date', 'client_id']
    for col in required_cols:
        if col not in col_map:
            print(f"WARNING: Column '{col}' not found in Client_Balances")
    
    # Reset all balances to 0 first; the sheet is the source of truth.
    # Clients absent from the sheet correctly become 0.
    cursor.execute("UPDATE clients SET outstandingBalance = 0, oldestUnpaidDate = NULL")
    updated = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), 2):
        try:
            client_name_cell = row[col_map.get('client_name', 1) - 1]
            outstanding_cell = row[col_map.get('outstanding_balance', 2) - 1]
            oldest_date_cell = row[col_map.get('oldest_unpaid_invoice_date', 3) - 1]
            excel_client_id_cell = row[col_map.get('client_id', 4) - 1]
            
            client_name = client_name_cell.value
            outstanding = outstanding_cell.value if outstanding_cell else 0
            oldest_date = oldest_date_cell.value
            excel_client_id = excel_client_id_cell.value if excel_client_id_cell else None
            
            if not client_name:
                continue
            
            # Convert outstanding to float
            try:
                outstanding = float(outstanding) if outstanding else 0
            except:
                outstanding = 0
            
            # Parse date
            oldest_date_parsed = parse_date(oldest_date)
            
            # Client ID == Excel ID by design (post-migration). Look up by id directly.
            if excel_client_id:
                cursor.execute("SELECT id FROM clients WHERE id = %s", (excel_client_id,))
                result = cursor.fetchone()
                if result:
                    client_id = result[0]
                else:
                    # Defensive fallback: fuzzy name match (should not happen post-migration).
                    client_id = find_client_by_name(cursor, str(client_name))
            else:
                client_id = find_client_by_name(cursor, str(client_name))
            
            if not client_id:
                print(f"  SKIP: Client '{client_name}' (Excel ID: {excel_client_id}) not found")
                continue
            
            # Update client with balance data; keep excelClientId == id for self-healing.
            cursor.execute("""
                UPDATE clients 
                SET outstandingBalance = %s, 
                    oldestUnpaidDate = %s,
                    excelClientId = id
                WHERE id = %s
            """, (outstanding, oldest_date_parsed, client_id))
            
            updated += 1
            
        except Exception as e:
            print(f"  ERROR on row {row_idx}: {e}")
            continue
    
    return updated

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 auto-sync.py <excel_file>")
        sys.exit(1)
    
    excel_file = sys.argv[1]
    
    if not os.path.exists(excel_file):
        print(f"ERROR: File not found: {excel_file}")
        sys.exit(1)
    
    print(f"[{datetime.now().isoformat()}] Starting sync from {excel_file}")
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        # Sync Client_Balances
        print("Syncing Client_Balances...")
        updated = sync_client_balances(cursor, excel_file)
        print(f"  ✓ Updated {updated} clients with balance data")
        
        conn.commit()
        cursor.close()
        conn.close()
        
        print(f"[{datetime.now().isoformat()}] Sync completed successfully")
        
    except Exception as e:
        print(f"ERROR: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()

