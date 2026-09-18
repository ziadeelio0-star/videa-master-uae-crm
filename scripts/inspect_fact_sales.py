"""Inspect Fact_Sales sheet header and rows."""
from openpyxl import load_workbook
XLSX = "/home/ubuntu/upload/Videa-Master-DashboardFinale.xlsx"
wb = load_workbook(XLSX, data_only=True, read_only=True)
print("sheets:", wb.sheetnames)

# Pick the sales sheet
sname = None
for s in wb.sheetnames:
    if 'sale' in s.lower() or 'fact' in s.lower():
        sname = s
        break
print("Using sheet:", sname)

ws = wb[sname]
rows = list(ws.iter_rows(values_only=True))
print("total rows:", len(rows))
print("header:", rows[0])
print("sample rows:")
for r in rows[1:8]:
    print(r)
print("\ncolumn H samples (Item_ID?):")
for r in rows[1:15]:
    if len(r) >= 8:
        print(repr(r[7]))
